import sys
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

# See test_applicant.py: the repo root __init__.py makes pytest resolve this
# module as UBJob_Application_Agent.test_generator_docx, so the parent of the
# repo (not the repo itself) is what needs to be on sys.path.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from UBJob_Application_Agent import generator  # noqa: E402

JOB = {
    "Job_Title": "Data Analyst",
    "Job_ID": "R-12345",
    "Department": "Institutional Research",
    "Campus": "North Campus",
    "Description": "Analyze enrollment data with SQL and Python.",
}
PERSON = {
    "first_name": "Ada",
    "last_name": "Lovelace",
    "email": "ada@example.com",
    "phone": "555-0100",
    "location": "Buffalo, NY",
    "linkedin": "linkedin.com/in/ada",
    "github": "github.com/ada",
}


def _text(path):
    return [p.text for p in Document(path).paragraphs]


class TestSanitizeFilename:
    def test_unsafe_characters_become_single_underscores(self):
        assert generator._sanitize_filename("R 12/34: x") == "R_12_34_x"

    def test_empty_and_none_become_unknown(self):
        assert generator._sanitize_filename("") == "Unknown"
        assert generator._sanitize_filename(None) == "Unknown"

    def test_path_traversal_cannot_escape_the_output_dir(self):
        name = generator._sanitize_filename("../../etc/passwd")
        assert "/" not in name and "\\" not in name and ".." not in name

    def test_dashes_and_underscores_survive(self):
        assert generator._sanitize_filename("R-1_2") == "R-1_2"


class TestSkillsHighlight:
    @pytest.mark.parametrize("kind, word", [
        ("DATA", "Python"),
        ("research", "experimental design"),
        ("Associate", "coordination"),
    ])
    def test_known_resume_types(self, kind, word):
        assert word in generator._skills_highlight_by_resume_type(kind)

    @pytest.mark.parametrize("kind", ["", None, "UNKNOWN"])
    def test_unknown_or_missing_type_gets_the_generic_hint(self, kind):
        assert generator._skills_highlight_by_resume_type(kind) == (
            "Highlight the most relevant strengths for this role."
        )


class TestAddBodyParagraphs:
    def test_dash_lines_become_bullets_and_others_plain_paragraphs(self):
        doc = Document()
        generator._add_body_paragraphs(doc, "Intro line\n- first\n- second\n\nClosing line")
        got = [(p.text, p.style.name) for p in doc.paragraphs]
        assert got == [
            ("Intro line", "Normal"),
            ("first", "List Bullet"),
            ("second", "List Bullet"),
            ("Closing line", "Normal"),
        ]

    def test_trailing_bullets_are_flushed(self):
        doc = Document()
        generator._add_body_paragraphs(doc, "- only\n- bullets")
        assert [p.style.name for p in doc.paragraphs] == ["List Bullet", "List Bullet"]

    def test_blank_lines_and_surrounding_whitespace_are_ignored(self):
        doc = Document()
        generator._add_body_paragraphs(doc, "\n\n  padded  \n\n\n")
        assert [p.text for p in doc.paragraphs] == ["padded"]


class TestReplacePlaceholders:
    def test_paragraphs_and_table_cells_are_replaced(self):
        doc = Document()
        doc.add_paragraph("Dear {{DEPARTMENT}},")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Re: {{TITLE}}"
        assert generator._replace_placeholders(doc, {"{{DEPARTMENT}}": "HR", "{{TITLE}}": "Analyst"})
        assert doc.paragraphs[0].text == "Dear HR,"
        assert table.cell(0, 0).text == "Re: Analyst"

    def test_returns_false_when_no_placeholder_is_present(self):
        doc = Document()
        doc.add_paragraph("Nothing to replace")
        assert generator._replace_placeholders(doc, {"{{X}}": "y"}) is False


class TestGenerateDocx:
    def test_fallback_letter_without_a_template(self, tmp_path):
        path = generator.generate_cover_letter(
            JOB, "DATA", PERSON, output_dir=str(tmp_path), template_path=None, use_llm=False
        )
        assert Path(path).name == "Cover_Letter_R-12345.docx"
        text = _text(path)
        assert text[0] == "Ada Lovelace"
        assert "Buffalo, NY | 555-0100" in text
        assert "Institutional Research" in text
        assert "North Campus" in text
        assert "Dear Hiring Committee," in text
        assert text[-2:] == ["Sincerely,", "Ada Lovelace"]
        assert any("Data Analyst role in Institutional Research" in t for t in text)

    def test_output_dir_is_created(self, tmp_path):
        out = tmp_path / "nested" / "docs"
        generator.generate_cover_letter(JOB, "DATA", PERSON, output_dir=str(out), template_path=None, use_llm=False)
        assert out.is_dir()

    def test_missing_optional_fields_do_not_raise(self, tmp_path):
        path = generator.generate_cover_letter(
            {"Job_ID": "X1"}, "", {}, output_dir=str(tmp_path), template_path=None, use_llm=False
        )
        assert Path(path).exists()

    def test_template_placeholders_are_filled(self, tmp_path):
        template = Document()
        template.add_paragraph("{{FIRST_NAME}} {{LAST_NAME}} | {{EMAIL}}")
        template.add_paragraph("{{DATE}}")
        template.add_paragraph("Re: {{TITLE}} ({{JOB_ID}}) at {{DEPARTMENT}}")
        template.add_paragraph("{{BODY}}")
        tpl_path = tmp_path / "tpl.docx"
        template.save(str(tpl_path))

        path = generator.generate_cover_letter(
            JOB, "DATA", PERSON, output_dir=str(tmp_path / "out"),
            template_path=str(tpl_path), use_llm=False,
        )
        text = _text(path)
        assert text[0] == "Ada Lovelace | ada@example.com"
        assert text[2] == "Re: Data Analyst (R-12345) at Institutional Research"
        assert "{{" not in "\n".join(text)
        assert "I am excited to apply for the Data Analyst role" in text[3]

    def test_template_without_placeholders_falls_back_to_the_built_in_layout(self, tmp_path):
        template = Document()
        template.add_paragraph("A template that has no placeholders at all.")
        tpl_path = tmp_path / "plain.docx"
        template.save(str(tpl_path))

        path = generator.generate_cover_letter(
            JOB, "DATA", PERSON, output_dir=str(tmp_path / "out"),
            template_path=str(tpl_path), use_llm=False,
        )
        text = _text(path)
        assert "A template that has no placeholders at all." not in text
        assert "Dear Hiring Committee," in text

    def test_missing_template_file_uses_the_built_in_layout(self, tmp_path):
        path = generator.generate_cover_letter(
            JOB, "DATA", PERSON, output_dir=str(tmp_path),
            template_path=str(tmp_path / "does_not_exist.docx"), use_llm=False,
        )
        assert "Dear Hiring Committee," in _text(path)

    def test_job_id_with_unsafe_characters_is_sanitized_in_the_filename(self, tmp_path):
        job = dict(JOB, Job_ID="R/12:34")
        path = generator.generate_cover_letter(job, "DATA", PERSON, output_dir=str(tmp_path),
                                               template_path=None, use_llm=False)
        assert Path(path).parent == tmp_path
        assert Path(path).name == "Cover_Letter_R_12_34.docx"


class _FakeResponses:
    def __init__(self, text):
        self.text = text
        self.kwargs = None

    def create(self, **kwargs):
        self.kwargs = kwargs
        return SimpleNamespace(output_text=self.text)


@pytest.fixture
def fake_client(monkeypatch):
    responses = _FakeResponses("  Generated body.\n- point one  \n")
    monkeypatch.setattr(generator, "_get_client", lambda: SimpleNamespace(responses=responses))
    return responses


class TestLlmBody:
    def test_output_is_stripped(self, fake_client):
        body = generator.generate_cover_letter_body_llm(JOB, "DATA", PERSON)
        assert body == "Generated body.\n- point one"

    def test_prompt_contains_the_job_candidate_and_resume_guidance(self, fake_client):
        generator.generate_cover_letter_body_llm(JOB, "RESEARCH", PERSON)
        user = fake_client.kwargs["input"][1]["content"]
        assert "Ada Lovelace" in user
        assert "Data Analyst" in user and "R-12345" in user and "Institutional Research" in user
        assert "Analyze enrollment data" in user
        assert "experimental design" in user  # RESEARCH guidance
        assert fake_client.kwargs["input"][0]["role"] == "system"

    def test_custom_tone_and_template_prompt_replace_the_defaults(self, fake_client):
        generator.generate_cover_letter_body_llm(
            JOB, "DATA", PERSON, tone_instructions="Be terse.", template_prompt="Use my outline."
        )
        user = fake_client.kwargs["input"][1]["content"]
        assert "Be terse." in user and "Use my outline." in user
        assert "3–5 short paragraphs" not in user

    def test_no_template_prompt_says_none(self, fake_client):
        generator.generate_cover_letter_body_llm(JOB, "DATA", PERSON)
        assert fake_client.kwargs["input"][1]["content"].rstrip().endswith("None")

    def test_missing_job_fields_fall_back_to_generic_wording(self, fake_client):
        generator.generate_cover_letter_body_llm({}, "DATA", {})
        user = fake_client.kwargs["input"][1]["content"]
        assert "Title: the role" in user and "Department/Unit: the department" in user

    def test_docx_uses_the_llm_body_with_bullets(self, fake_client, tmp_path):
        path = generator.generate_cover_letter(
            JOB, "DATA", PERSON, output_dir=str(tmp_path), template_path=None, use_llm=True
        )
        doc = Document(path)
        assert "Generated body." in [p.text for p in doc.paragraphs]
        bullets = [p.text for p in doc.paragraphs if p.style.name == "List Bullet"]
        assert bullets == ["point one"]
